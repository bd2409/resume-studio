"""Three-step tailoring pipeline: extract a fact bank, write the tailored resume, verify it.

Every step returns validated JSON (structured outputs), and the writer must cite the facts
behind each bullet so the verifier can reject anything not traceable to the customer's resumes.
"""
import base64
import json
import subprocess
import tempfile
from pathlib import Path

import anthropic
from docx import Document

from . import config, render
from .schemas import CoverLetter, LinkedInProfile, Draft, Profile, Resume, Verified

_client = None


def client():
    global _client
    if _client is None:
        _client = anthropic.Anthropic(api_key=config.ANTHROPIC_API_KEY, timeout=600, max_retries=3)
    return _client


# ---------------------------------------------------------------- reading uploads

def _docx_text(path):
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(dict.fromkeys(cells)))
    return "\n".join(parts)


def _convert_with_libreoffice(path, fmt):
    out_dir = tempfile.mkdtemp()
    with render._soffice_lock:
        subprocess.run(
            ["soffice", "--headless", "--convert-to", fmt, "--outdir", out_dir, str(path)],
            check=True, capture_output=True, timeout=120,
        )
    return next(Path(out_dir).glob(f"*.{fmt}"))


def resume_blocks(paths):
    """Turn uploaded files into Claude content blocks: PDFs as documents, Word files as text."""
    blocks = []
    for i, path in enumerate(paths, 1):
        path = Path(path)
        ext = path.suffix.lower()
        if ext in {".doc", ".rtf", ".odt"}:
            path, ext = _convert_with_libreoffice(path, "docx"), ".docx"
        if ext == ".pdf":
            data = base64.standard_b64encode(path.read_bytes()).decode()
            blocks.append({
                "type": "document",
                "source": {"type": "base64", "media_type": "application/pdf", "data": data},
                "title": f"Resume {i}",
            })
        elif ext == ".docx":
            blocks.append({"type": "text", "text": f"<resume index=\"{i}\">\n{_docx_text(path)}\n</resume>"})
        else:
            text = path.read_text(encoding="utf-8", errors="ignore")
            blocks.append({"type": "text", "text": f"<resume index=\"{i}\">\n{text}\n</resume>"})
    return blocks


# ---------------------------------------------------------------- prompts

EXTRACT_SYSTEM = """You are a meticulous resume analyst. You receive up to six resumes written by the same person over time. Build one complete, accurate fact bank of their career.

Rules:
- Record only what the resumes actually say. Never infer, embellish or round numbers.
- Merge the same job across resumes into one role. Keep the most complete and most recent wording of each achievement, and keep every distinct achievement.
- Keep every number, currency figure, percentage, team size, client name and tool exactly as written.
- Give each fact a short unique id (F1, F2, ...) across the whole profile.
- Header: take the most recent contact details. If a field is missing, use an empty string. The headline is the person's most recent title plus their core domain, in under 12 words.
- List any contradictions between resumes (dates, titles, figures) in conflicts."""

WRITE_SYSTEM = """You are a senior resume writer for Indian and global job markets. Customers pay for a resume that clearly beats what they could write themselves, so every line must earn its place.

You receive the candidate's fact bank (from their own resumes) and a target job description. Write a resume tailored to that job, in this fixed structure: header, PROFESSIONAL SUMMARY, KEY COMPETENCIES, PROFESSIONAL EXPERIENCE, EDUCATION, CERTIFICATIONS & RECOGNITIONS, ACHIEVEMENTS.

First, analyse the job: the target role, its must-have skills and the exact keywords an applicant tracking system will scan for.

Truthfulness (non-negotiable):
- Use only facts from the fact bank. Every experience bullet lists the fact ids it is based on.
- Never invent employers, titles, dates, degrees, certifications, tools, metrics, team sizes or results. Never change a number.
- You may reword, combine and reframe facts to show their relevance, and use the job's terminology where the underlying experience genuinely matches.
- If the job needs something the facts do not support, leave it out and add a gap note advising the customer, for example: "The role asks for Power BI. Your resumes don't mention it; add it only if you have used it."

Quality:
- Headline: target-role-aligned title plus 2 to 3 strengths, under 12 words, true to the candidate's level.
- Summary: 3 to 4 sentences, about 60 to 80 words. Years of experience, domain, 2 to 3 strongest quantified results relevant to this job, and what they bring to this role. No first person, no cliches ("results-driven", "dynamic", "go-getter", "passionate").
- Key competencies: 9 to 12 short items (2 to 4 words), ordered by relevance to the job, each backed by the fact bank, using the job's wording where accurate.
- Experience: most recent first. 4 to 6 bullets for the latest one or two relevant roles, 2 to 3 for older ones, 1 to 2 for roles over 10 years old. Order bullets by relevance to the job.
- Bullets: start with a strong, varied past-tense action verb (present tense for the current role), show scope, action and measurable result, one to two lines each. Where a bullet has a clear theme, start it with a short bold-style lead phrase followed by a colon, for example "Recurring payments: Launched ...". Use lead phrases on the main role's bullets only.
- Context: one short line on the company if known from the facts, else empty.
- Dates as "Mon YYYY - Mon YYYY" or "Mon YYYY - Present".
- Education: highest first, with grade if given.
- Certifications & recognitions: certifications and formal awards. Achievements: other standout accomplishments (awards, rankings, publications, talks, patents, notable projects) not already used as bullets. If there are none, use the strongest remaining facts rather than repeating bullets.
- Length: fits one page for under 5 years of experience, two pages at most otherwise.
- Indian conventions are fine (Rs, lakh, crore) when the source uses them. Use plain text only, no markdown."""

VERIFY_SYSTEM = """You are the quality checker for a paid resume-writing service. You receive the candidate's fact bank, the job description and a draft resume whose bullets cite fact ids. Return the final, corrected resume.

Check every line:
1. Truth: each claim, number, tool, title, date, employer and qualification must be supported by the fact bank. Remove or rewrite anything unsupported, including in the summary, headline and competencies. Numbers must match the source exactly.
2. Fit: the strongest job-relevant facts are used and ordered first; must-have keywords appear wherever they are honestly supported.
3. Writing: no repeated opening verbs within a role, no cliches, no first person, no filler, consistent tense and date format, no spelling or grammar errors, plain text only.
4. Format: sections complete; bullet counts and length follow the rules (one page under 5 years of experience, at most two pages otherwise).

List each correction you made. Keep the gap notes that still apply and add any new ones. Do not add anything that is not in the fact bank."""


COVER_SYSTEM = """You write cover letters for a paid resume service. You receive the candidate's fact bank, the job description (or just a target role) and their final tailored resume. Write a cover letter for this application.

Rules:
- 250 to 350 words in 3 or 4 paragraphs: why this role and company (only what the job description says about them), two or three of the most relevant quantified achievements, how they fit the must-have requirements, and a short confident close.
- Use only facts from the fact bank and the resume. Never invent experience, numbers, tools or knowledge of the company that is not in the job description.
- Specific, warm and professional. No cliches ("I am writing to express", "results-driven", "passionate"), no flattery, no repeating the resume line by line.
- Address the company by name only if the job description names it; otherwise use "Dear Hiring Manager,".
- Plain text only, no placeholders such as [Company]."""


LINKEDIN_SYSTEM = """You are a LinkedIn profile strategist who has helped thousands of professionals get found by recruiters. You receive the person's fact bank (from their own resumes) and the role they are targeting. Rewrite their LinkedIn profile.

Truthfulness (non-negotiable): use only facts from the fact bank. Never invent employers, titles, dates, numbers, tools or results.

Headline (max 220 characters): target role and level, 2 to 3 specialisms recruiters search for, and one proof point if it fits. Use " | " separators. No emojis, no "aspiring", no "open to work".
About (1,500 to 2,500 characters): written in first person, warm and confident. Open with a hook line about the value they create, then their story and focus, 3 to 5 quantified wins (can be a short list with "-" bullets), what they are looking for or how they help, and a closing line inviting contact. Weave in the keywords recruiters for the target role search for. Short paragraphs separated by blank lines.
Experience: for each role in the last 10 years, a description of 600 to 1,500 characters: one line of scope, then 3 to 5 achievement lines starting with "-". First person implied, no "I" at the start of every line.
Skills: 15 to 25, ordered by importance for the target role, each supported by the fact bank, using LinkedIn's standard skill names where possible.
Featured ideas, banner tagline and tips: specific to this person, practical, based on their real work.
Plain text only (no markdown asterisks or headings). Indian conventions (Rs, lakh, crore) are fine where the source uses them."""


def job_brief(job_description: str, target_role: str = ""):
    """The job text the writer works from. Without a posting, the writer optimises for the target role in general."""
    if job_description.strip():
        return job_description
    role = target_role.strip() or "the candidate's current career direction"
    return (f"No specific job posting was provided. Target role: {role}. "
            "Improve the resume for this kind of role in general: infer the skills and keywords that typical postings "
            "for it in India require, and present the candidate's real experience as strongly as possible for it.")


def match_score(must_have, covered):
    """Share of the job's must-have keywords that the final resume covers."""
    must = {k.strip().lower() for k in must_have if k.strip()}
    have = {k.strip().lower() for k in covered}
    if not must:
        return None
    hit = must & have
    return {"covered": len(hit), "total": len(must), "percent": round(100 * len(hit) / len(must)),
            "missing": sorted(k for k in must_have if k.strip().lower() in must - hit)}


def _call(system, content, schema, effort):
    response = client().messages.parse(
        model=config.AI_MODEL,
        max_tokens=16000,
        system=system,
        output_config={"effort": effort},
        messages=[{"role": "user", "content": content}],
        output_format=schema,
    )
    if response.stop_reason == "refusal":
        raise RuntimeError("The AI declined this request.")
    if response.stop_reason == "max_tokens" or response.parsed_output is None:
        raise RuntimeError("The AI response was incomplete.")
    return response.parsed_output


def extract_profile(paths):
    content = resume_blocks(paths) + [{
        "type": "text",
        "text": "Build the fact bank from the resumes above.",
    }]
    return _call(EXTRACT_SYSTEM, content, Profile, effort="medium")


def write_and_verify(profile: Profile, job_description: str, instructions: str = ""):
    """Returns (Resume, notes) where notes holds gap notes and corrections for the customer."""
    facts = profile.model_dump_json(indent=1)
    extra = f"\n\n<customer_instructions>\n{instructions}\n</customer_instructions>" if instructions else ""
    draft = _call(
        WRITE_SYSTEM,
        f"<fact_bank>\n{facts}\n</fact_bank>\n\n<job_description>\n{job_description}\n</job_description>{extra}",
        Draft,
        effort="high",
    )
    verified = _call(
        VERIFY_SYSTEM,
        f"<fact_bank>\n{facts}\n</fact_bank>\n\n<job_description>\n{job_description}\n</job_description>\n\n"
        f"<draft>\n{draft.model_dump_json(indent=1)}\n</draft>",
        Verified,
        effort="high",
    )
    notes = {
        "target_role": draft.target_role,
        "gap_notes": verified.gap_notes,
        "corrections": verified.corrections,
        "keyword_coverage": verified.keyword_coverage,
        "match": match_score(draft.must_have_keywords, verified.keyword_coverage),
        "conflicts": profile.conflicts,
    }
    return verified.resume, notes


def write_linkedin(profile: Profile, job: str, instructions: str = ""):
    extra = f"\n\n<customer_instructions>\n{instructions}\n</customer_instructions>" if instructions else ""
    return _call(
        LINKEDIN_SYSTEM,
        f"<fact_bank>\n{profile.model_dump_json(indent=1)}\n</fact_bank>\n\n<target>\n{job}\n</target>{extra}",
        LinkedInProfile,
        effort="high",
    )


def mock_linkedin():
    return LinkedInProfile(
        headline="Senior Product Manager | Payments & Fintech Platforms | Scaled UPI AutoPay to Rs 1,200 Cr | B2B SaaS",
        about=("This is sample content, shown because the AI is not connected yet.\n\n"
               "With the AI connected, this About section is written from the customer's real resumes, in first person, "
               "with their strongest quantified wins and the keywords recruiters search for."),
        experience=[{"company": "PayNova Technologies", "title": "Senior Product Manager",
                     "description": "Own merchant payments for 2M+ small businesses.\n- Launched UPI AutoPay and e-mandate flows, adding Rs 1,200 Cr in annual processed volume.\n- Cut merchant activation time from 5 days to 36 hours."}],
        skills=["Product Management", "Payments", "Fintech", "B2B SaaS", "Product Strategy", "SQL"],
        featured_ideas=["A short post on how UPI AutoPay was launched", "The Global Fintech Fest 2024 talk"],
        banner_tagline="Building payments that small businesses love",
        tips=["Add a professional photo with a plain background", "Turn on Creator mode and post weekly about payments"],
    )


def write_cover_letter(profile: Profile, job_description: str, resume: Resume):
    return _call(
        COVER_SYSTEM,
        f"<fact_bank>\n{profile.model_dump_json(indent=1)}\n</fact_bank>\n\n<job_description>\n{job_description}\n</job_description>\n\n"
        f"<resume>\n{resume.model_dump_json(indent=1)}\n</resume>",
        CoverLetter,
        effort="medium",
    )


# ---------------------------------------------------------------- mock mode for local testing

SAMPLE = Path(__file__).with_name("sample_resume.json")


def mock_generate():
    resume = Resume.model_validate(json.loads(SAMPLE.read_text()))
    notes = {
        "target_role": "Sample output (AI not connected)",
        "gap_notes": ["This is sample content. Set ANTHROPIC_API_KEY to generate real resumes."],
        "corrections": [],
        "keyword_coverage": [],
        "match": {"covered": 13, "total": 15, "percent": 87, "missing": ["Power BI", "Kafka"]},
        "conflicts": [],
    }
    return resume, notes


def mock_cover_letter(resume: Resume):
    return CoverLetter(
        greeting="Dear Hiring Manager,",
        paragraphs=[
            "This is a sample cover letter, shown because the AI is not connected yet.",
            f"With the AI connected, this letter is written from {resume.header.name}'s real experience and the job description, in 250 to 350 words.",
            "It highlights the most relevant achievements and closes with a short, confident note.",
        ],
        closing="Sincerely,",
    )
