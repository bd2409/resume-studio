"""Render resume content into the fixed Word template, and convert it to PDF."""
import json
import subprocess
import sys
import threading
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

GREY = RGBColor(0x44, 0x44, 0x44)
RIGHT_TAB = Cm(18.0)  # usable width on A4 with 1.5 cm margins

# Every template keeps the same sections and a single ATS-friendly column; only the look changes.
# Fonts are ones LibreOffice has metric-compatible copies of (Carlito, Caladea, Liberation), so the PDF matches Word.
TEMPLATES = {
    "classic": {"label": "Classic", "font": "Calibri", "accent": "1F3A5F", "align": "center", "titles": "rule", "name_size": 20},
    "executive": {"label": "Executive", "font": "Cambria", "accent": "6B1D2A", "align": "center", "titles": "double", "name_size": 21},
    "modern": {"label": "Modern", "font": "Calibri", "accent": "0F766E", "align": "left", "titles": "band", "name_size": 22},
    "minimal": {"label": "Minimal", "font": "Arial", "accent": "222222", "align": "left", "titles": "plain", "name_size": 18},
}
DEFAULT_TEMPLATE = "classic"
_lock = threading.Lock()  # the active theme is module state, so builds run one at a time

FONT = "Calibri"
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)
_theme = TEMPLATES[DEFAULT_TEMPLATE]


def use_template(key):
    global FONT, ACCENT, _theme
    _theme = TEMPLATES.get(key) or TEMPLATES[DEFAULT_TEMPLATE]
    FONT = _theme["font"]
    ACCENT = RGBColor.from_string(_theme["accent"])


def set_base_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(2)
    pf.line_spacing = 1.08
    for sec in doc.sections:
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        sec.top_margin = sec.bottom_margin = Cm(1.4)
        sec.left_margin = sec.right_margin = Cm(1.5)


def bottom_border(paragraph, color=None, size=8, style="single"):
    color = color or _theme["accent"]
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), style)
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    pPr.append(pbdr)


def shade(paragraph, fill):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def run(p, text, bold=False, italic=False, size=None, color=None):
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    return r


def header(doc, h):
    align = WD_ALIGN_PARAGRAPH.CENTER if _theme["align"] == "center" else WD_ALIGN_PARAGRAPH.LEFT
    p = doc.add_paragraph()
    p.alignment = align
    name = h["name"] if _theme["titles"] == "band" else h["name"].upper()
    run(p, name, bold=True, size=_theme["name_size"], color=ACCENT)
    p.paragraph_format.space_after = Pt(0)

    p = doc.add_paragraph()
    p.alignment = align
    run(p, h["headline"], bold=True, size=11.5, color=GREY)
    p.paragraph_format.space_after = Pt(2)

    p = doc.add_paragraph()
    p.alignment = align
    parts = [h["email"], h["phone"], h["linkedin"], h["location"]]
    run(p, "  |  ".join(x for x in parts if x), size=9.5)
    if _theme["titles"] == "double":
        bottom_border(p, size=6, style="double")
    elif _theme["titles"] != "plain":
        bottom_border(p, size=12)
    else:
        bottom_border(p, color="BBBBBB", size=4)
    p.paragraph_format.space_after = Pt(4)


def section_title(doc, title):
    p = doc.add_paragraph()
    style = _theme["titles"]
    if style == "band":
        shade(p, _theme["accent"])
        run(p, " " + title.upper(), bold=True, size=11, color=RGBColor(0xFF, 0xFF, 0xFF))
    elif style == "plain":
        run(p, title.upper(), bold=True, size=11, color=ACCENT)
    else:
        run(p, title.upper(), bold=True, size=11.5, color=ACCENT)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    if style == "rule":
        bottom_border(p, size=6)
    elif style == "double":
        bottom_border(p, size=4, style="double")


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.55)
    p.paragraph_format.first_line_indent = Cm(-0.4)
    p.paragraph_format.space_after = Pt(1.5)
    # "Lead phrase: rest" renders the lead phrase in bold
    if ": " in text and len(text.split(": ", 1)[0]) <= 45:
        lead, rest = text.split(": ", 1)
        run(p, lead + ": ", bold=True)
        run(p, rest)
    else:
        run(p, text)


def left_right(doc, left, right, bold_left=True, italic=False, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.tab_stops.add_tab_stop(RIGHT_TAB, WD_TAB_ALIGNMENT.RIGHT)
    run(p, left, bold=bold_left, italic=italic, color=color)
    if right:
        run(p, "\t" + right, bold=bold_left, italic=italic, color=color)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    return p


def competencies(doc, items, cols=3):
    rows = (len(items) + cols - 1) // cols
    t = doc.add_table(rows=rows, cols=cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, item in enumerate(items):
        cell = t.cell(i % rows, i // rows)
        p = cell.paragraphs[0]
        run(p, "▪ ", color=ACCENT)
        run(p, item)
        p.paragraph_format.space_after = Pt(1)


def build(data, out, template=DEFAULT_TEMPLATE):
    with _lock:
        use_template(template)
        _build(data, out)


def _build(data, out):
    doc = Document()
    set_base_styles(doc)
    header(doc, data["header"])

    section_title(doc, "Professional Summary")
    p = doc.add_paragraph(data["summary"])
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    section_title(doc, "Key Competencies")
    competencies(doc, data["competencies"])

    section_title(doc, "Professional Experience")
    for i, job in enumerate(data["experience"]):
        p = left_right(doc, job["company"] + (", " + job["location"] if job.get("location") else ""), job["dates"])
        if i:
            p.paragraph_format.space_before = Pt(6)
        left_right(doc, job["title"], "", bold_left=False, italic=True, color=ACCENT)
        if job.get("context"):
            p = doc.add_paragraph()
            run(p, job["context"], italic=True, size=9.5, color=GREY)
            p.paragraph_format.space_after = Pt(1)
        for b in job["bullets"]:
            bullet(doc, b)

    section_title(doc, "Education")
    for ed in data["education"]:
        left_right(doc, ed["degree"], ed["dates"])
        p = doc.add_paragraph()
        run(p, ed["institution"] + (f"  |  {ed['grade']}" if ed.get("grade") else ""))
        p.paragraph_format.space_after = Pt(3)

    if data["certifications"]:
        section_title(doc, "Certifications & Recognitions")
        for c in data["certifications"]:
            bullet(doc, c)

    if data["achievements"]:
        section_title(doc, "Achievements")
        for a in data["achievements"]:
            bullet(doc, a)

    doc.core_properties.title = f"{data['header']['name']} - Resume"
    doc.core_properties.author = data["header"]["name"]
    doc.save(out)


def to_pdf(docx_path):
    """Convert with LibreOffice so the PDF matches the Word file exactly."""
    docx_path = Path(docx_path)
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(docx_path.parent), str(docx_path)],
        check=True, capture_output=True, timeout=180,
    )
    return docx_path.with_suffix(".pdf")


def to_previews(pdf_path, dpi=110):
    """Page images for the review screen; PDFs in iframes don't display on most phones."""
    pdf_path = Path(pdf_path)
    for old in pdf_path.parent.glob("preview-*.png"):
        old.unlink()
    subprocess.run(
        ["pdftoppm", "-r", str(dpi), "-png", str(pdf_path), str(pdf_path.parent / "preview")],
        check=True, capture_output=True, timeout=60,
    )
    return sorted(pdf_path.parent.glob("preview-*.png"))


if __name__ == "__main__":
    with open(sys.argv[1], encoding="utf-8") as f:
        build(json.load(f), sys.argv[2], *(sys.argv[3:4]))
